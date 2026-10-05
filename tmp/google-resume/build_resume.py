from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT

ROOT = Path('/Users/rafaelcarbajal/Desktop/Rafael_Portfolio')
OUT = ROOT / 'output/docx/Rafael_Carbajal_Google_UX_Intern_2027.docx'
OUT.parent.mkdir(parents=True, exist_ok=True)
doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Inches(8.5), Inches(11)
sec.top_margin = sec.bottom_margin = Inches(.5)
sec.left_margin = sec.right_margin = Inches(.6)
for name in ['Normal', 'Title', 'Heading 1', 'Heading 2', 'List Bullet']:
    s = doc.styles[name]
    s.font.name = 'Arial'
    s.font.color.rgb = RGBColor(0,0,0)
    s.font.size = Pt(10.5)
    s.paragraph_format.space_after = Pt(0)
    s.paragraph_format.space_before = Pt(0)
    s.paragraph_format.line_spacing = 1.06
normal = doc.styles['Normal']
title = doc.styles['Title']
title.font.size = Pt(20)
title.font.bold = True
title.paragraph_format.space_after = Pt(3)
h = doc.styles['Heading 1']
h.font.size = Pt(10.5)
h.font.bold = True
h.paragraph_format.space_before = Pt(7)
h.paragraph_format.space_after = Pt(4)
h.paragraph_format.keep_with_next = True

def para(text='', size=None, after=0, bold=False):
    p = doc.add_paragraph()
    r = p.add_run(text)
    r.bold = bold
    if size: r.font.size = Pt(size)
    p.paragraph_format.space_after = Pt(after)
    return p

def link(p, label, url):
    el=OxmlElement('w:hyperlink')
    el.set(qn('r:id'),p.part.relate_to(url,RT.HYPERLINK,is_external=True))
    r=OxmlElement('w:r');pr=OxmlElement('w:rPr')
    c=OxmlElement('w:color');c.set(qn('w:val'),'000000');pr.append(c)
    sz=OxmlElement('w:sz');sz.set(qn('w:val'),'19');pr.append(sz)
    u=OxmlElement('w:u');u.set(qn('w:val'),'single');pr.append(u)
    r.append(pr);t=OxmlElement('w:t');t.text=label;r.append(t);el.append(r);p._p.append(el)

def head(left,right):
    p=doc.add_paragraph()
    p.paragraph_format.keep_with_next=True
    p.paragraph_format.space_before=Pt(3)
    p.paragraph_format.tab_stops.add_tab_stop(Inches(7.3),WD_TAB_ALIGNMENT.RIGHT)
    p.add_run(left).bold=True
    p.add_run('\t'+right).font.size=Pt(9.5)
    return p

def bullet(text):
    p=doc.add_paragraph(style='List Bullet')
    p.paragraph_format.left_indent=Inches(.13)
    p.paragraph_format.first_line_indent=Inches(-.13)
    p.paragraph_format.space_after=Pt(2.5)
    p.add_run(text)
    return p

doc.add_paragraph('RAFAEL CARBAJAL SANCHEZ', 'Title')
p=para('Los Angeles, CA  |  (424) 436-9220  |  ',size=9.5,after=3)
link(p,'ra.carbajalsa@gmail.com','mailto:ra.carbajalsa@gmail.com')
p=doc.add_paragraph()
link(p,'rafaelsanchez.design','https://rafaelsanchez.design')
p.add_run('  |  ').font.size=Pt(9.5)
link(p,'linkedin.com/in/rafael-carbajal-956485285','https://www.linkedin.com/in/rafael-carbajal-956485285')

doc.add_paragraph('EDUCATION','Heading 1')
head('Santa Monica College | B.S. in Interaction Design','Expected 06/27')
para('Google UX Design Certificate, Coursera | 2025',size=9.5)

doc.add_paragraph('EXPERIENCE','Heading 1')
head('Snap Inc. | Snap Design Academy','Jun 2026 - Aug 2026')
para('Design Scholar (Contract) | Los Angeles, CA',size=9.5,after=3)
bullet('Led UX for Resources, a Snapchat local-support concept; aligned design requirements with engineering and marketing in a cross-functional sprint that produced a functional prototype in 3 weeks.')
bullet('Used 2 rounds of research and feedback with S.P.Y. and St. Joseph Center to scope the experience around service discovery and direct handoffs to organizations equipped to provide support.')
bullet('Designed entry points across Snap Map, Search, Spotlight, and sharing to fit existing user behaviors; created interfaces within Snapchat\u2019s visual language and presented the prototype to 300+ Snap employees.')

head('Plataforma Impact | UX Intern','Mar 2021 - Jan 2022')
para('Mexico City, Mexico (Remote)',size=9.5,after=3)
bullet('Created Figma layouts and implemented front-end updates for 3 client websites, collaborating with senior designers and engineers and supporting 1 website launch.')

doc.add_paragraph('SELECTED PROJECTS','Heading 1')
head('Pantri | Meal Planning and Digital Pantry','Spring 2026')
para('Lead UX / Product Designer | Team project',size=9.5,after=3)
bullet('Led UX and interface design across onboarding, meal planning, shopping, and cooking; used 13 in-store shopper interviews to shape receipt-based pantry updates that avoid manual item entry.')
bullet('Designed Figma user interface prototypes from low to high fidelity; built a 26-screen React prototype using AI-assisted coding and simulated data.')
bullet('Designed recovery paths for substitutions, quantity changes, and skipped meals after usability testing showed shoppers needed meal plans to adapt without restarting.')
bullet('Compared produce-sticker palettes in focus groups and tested with people with color vision deficiencies to refine ripeness cues within the sensing technology\u2019s color limits.')

head('HerWay | Community Transit App','Fall 2025')
para('Solo UX / UI Designer',size=9.5,after=3)
bullet('Used findings from 40 rider interviews to identify user needs and develop personas for women using public transit in South Los Angeles.')
bullet('Created 36 UI mockup screens, a 9-panel sketched journey storyboard, and an interactive Figma prototype linking route planning with community features.')
bullet('Revised navigation across 3 usability-testing rounds with 5 participants per round; brought routing and community into Home after testing showed riders overlooked features in a separate section.')

head('Nibble | Deployed AI Portfolio Assistant','Summer 2026')
para('Designer / Developer | Personal project',size=9.5,after=3)
bullet('Designed and built conversational navigation that connects visitor questions to case studies, with 4 visual states and streamed responses to communicate system progress.')
bullet('Launched an assistant that answered 400+ visitor questions in its first 4 weeks; Vercel reporting showed project pages viewed per visit rising from 1.6 to 2.3 over that period.')

doc.add_paragraph('SKILLS','Heading 1')
for label,text in [
    ('Design and research: ', 'User experience, interaction design, user flows, wireframing, prototyping, information architecture, accessibility, user interviews, usability testing, research synthesis'),
    ('Tools: ', 'Figma, Adobe Creative Cloud (Illustrator, Photoshop, InDesign)'),
    ('Development: ', 'React, JavaScript, HTML/CSS, AI-assisted coding')
]:
    p=doc.add_paragraph();p.add_run(label).bold=True;p.add_run(text)
    p.paragraph_format.space_after=Pt(2)

doc.core_properties.title='Rafael Carbajal Sanchez Resume'
doc.core_properties.subject='Google User Experience Design Intern Summer 2027'
doc.core_properties.author='Rafael Carbajal Sanchez'
for tree in [doc.styles.element, doc.element]:
    for border in tree.xpath('.//w:pBdr'):
        border.getparent().remove(border)
for level in doc.part.numbering_part.element.xpath('.//w:lvl'):
    fmt=level.find(qn('w:numFmt'))
    if fmt is not None and fmt.get(qn('w:val'))=='bullet':
        level.find(qn('w:lvlText')).set(qn('w:val'),'\u2022')
        props=level.find(qn('w:rPr'))
        if props is not None:
            fonts=props.find(qn('w:rFonts'))
            if fonts is not None:
                fonts.set(qn('w:ascii'),'Arial')
                fonts.set(qn('w:hAnsi'),'Arial')
doc.save(OUT)
print(OUT)
print('Words:',len(' '.join(p.text for p in doc.paragraphs).split()))
